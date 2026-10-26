import os
import re
from pathlib import Path

from dotenv import load_dotenv
from pypdf import PdfReader
from docx import Document

from langchain_core.documents import Document as LCDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter

from langchain_ollama import OllamaEmbeddings
from langchain_qdrant import QdrantVectorStore

load_dotenv()

DATA_DIR = Path("data")
OLLAMA_BASE_URL = os.getenv("OLLAMA_BASE_URL", "http://localhost:11434")
EMBEDDING_MODEL = os.getenv("EMBEDDING_MODEL", "nomic-embed-text")
QDRANT_URL = os.getenv("QDRANT_URL", "http://localhost:6333")
COLLECTION_NAME = os.getenv("QDRANT_COLLECTION", "my_collection")

# 1. Load PDF documents
def load_pdf(file_path: Path):
    reader = PdfReader(str(file_path))
    documents = []
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            documents.append(
                LCDocument(
                    page_content=text,
                    metadata={"source": file_path.name, "page": page_number}
                )
            )
    return documents

# 2. Load TXT documents
def load_txt(file_path: Path):
    text = file_path.read_text(encoding="utf-8", errors="ignore")
    return [
        LCDocument(
            page_content=text,
            metadata={"source": file_path.name, "page": 1}
        )
    ]

# 3. Load DOCX documents
def load_docx(file_path: Path):
    doc = Document(str(file_path))
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    return [
        LCDocument(
            page_content=text,
            metadata={"source": file_path.name, "page": 1}
        )
    ]

# 4. Load all supported documents
def load_documents():
    if not DATA_DIR.exists():
        raise FileNotFoundError("The data folder does not exist.")
    documents = []
    supported_extensions = {".pdf", ".txt", ".docx"}
    for file_path in DATA_DIR.iterdir():
        if not file_path.is_file():
            continue
        extension = file_path.suffix.lower()
        if extension not in supported_extensions:
            continue
        print(f"Loading: {file_path.name}")
        if extension == ".pdf":
            documents.extend(load_pdf(file_path))
        elif extension == ".txt":
            documents.extend(load_txt(file_path))
        elif extension == ".docx":
            documents.extend(load_docx(file_path))
    if not documents:
        raise ValueError("No supported documents found in data folder.")
    return documents

# 5. Clean text
def clean_documents(documents):
    cleaned_documents = []
    for document in documents:
        text = document.page_content
        text = re.sub(r"\s+", " ", text)
        text = text.strip()
        if text:
            cleaned_documents.append(
                LCDocument(
                    page_content=text,
                    metadata=document.metadata
                )
            )
    return cleaned_documents

# 6. Split documents into chunks
def split_documents(documents):
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=150,
        separators=["\n\n", "\n", ". ", " ", ""]
    )
    chunks = splitter.split_documents(documents)
    print(f"Total chunks: {len(chunks)}")
    return chunks

# 7. Create embedding model
def get_embedding_model():
    return OllamaEmbeddings(
        model=EMBEDDING_MODEL,
        base_url=OLLAMA_BASE_URL
    )

# 8. Store documents in Qdrant
def store_documents(chunks):
    embeddings = get_embedding_model()
    vector_store = QdrantVectorStore.from_documents(
        documents=chunks,
        embedding=embeddings,
        url=QDRANT_URL,
        collection_name=COLLECTION_NAME,
        force_recreate=True
    )
    print("Documents stored successfully in Qdrant!")
    return vector_store

# 9. Main indexing process
def ingest_documents():
    print("\n========== INDEXING STARTED ==========\n")
    documents = load_documents()
    print(f"Loaded documents: {len(documents)}")
    cleaned_documents = clean_documents(documents)
    print(f"Cleaned documents: {len(cleaned_documents)}")
    chunks = split_documents(cleaned_documents)
    store_documents(chunks)
    print("\n========== INDEXING COMPLETED ==========\n")

if __name__ == "__main__":
    ingest_documents()